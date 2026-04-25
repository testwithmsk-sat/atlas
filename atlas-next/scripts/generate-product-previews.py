from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader


CANVAS_SIZE = (1400, 1750)

CATEGORY_STYLES = {
    "Business": {
        "bg_top": (245, 238, 231),
        "bg_bottom": (232, 219, 207),
        "accent": (126, 89, 59),
        "accent_soft": (201, 168, 140),
        "panel": (255, 251, 245),
    },
    "Events & Parties": {
        "bg_top": (250, 242, 235),
        "bg_bottom": (240, 222, 214),
        "accent": (161, 95, 72),
        "accent_soft": (223, 177, 154),
        "panel": (255, 250, 246),
    },
    "Planners & Productivity": {
        "bg_top": (241, 243, 237),
        "bg_bottom": (223, 229, 214),
        "accent": (96, 112, 86),
        "accent_soft": (171, 188, 155),
        "panel": (250, 252, 248),
    },
}


def load_font(size: int, bold: bool = False, italic: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = []
    if bold and italic:
        candidates.extend(
            [
                r"C:\Windows\Fonts\georgiaz.ttf",
                r"C:\Windows\Fonts\segoeuiz.ttf",
            ]
        )
    elif bold:
        candidates.extend(
            [
                r"C:\Windows\Fonts\georgiab.ttf",
                r"C:\Windows\Fonts\segoeuib.ttf",
            ]
        )
    else:
        candidates.extend(
            [
                r"C:\Windows\Fonts\georgia.ttf",
                r"C:\Windows\Fonts\segoeui.ttf",
            ]
        )

    for candidate in candidates:
        if os.path.exists(candidate):
            return ImageFont.truetype(candidate, size=size)

    return ImageFont.load_default()


TITLE_FONT = load_font(80, bold=True)
TITLE_FONT_MEDIUM = load_font(70, bold=True)
TITLE_FONT_SMALL = load_font(60, bold=True)
SUBTITLE_FONT = load_font(34, bold=True)
BODY_FONT = load_font(30)
SMALL_FONT = load_font(24)
LABEL_FONT = load_font(26, bold=True)
MONO_FONT = load_font(24)


def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = []

    for word in words:
        trial = " ".join(current + [word]).strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current.append(word)
        else:
            if current:
                lines.append(" ".join(current))
            current = [word]

    if current:
        lines.append(" ".join(current))

    return lines


def extract_pdf_lines(file_path: str) -> list[str]:
    reader = PdfReader(file_path)
    if not reader.pages:
        return ["Digital PDF template", "Preview text was not available."]

    text = reader.pages[0].extract_text() or ""
    raw_lines = [normalize_spaces(line) for line in text.splitlines()]
    filtered = []

    for line in raw_lines:
        if not line:
            continue
        if line.lower().startswith("the digital atlas"):
            continue
        if line in filtered:
            continue
        filtered.append(line)
        if len(filtered) >= 8:
            break

    return filtered or ["Digital PDF template", "Preview text was not available."]


def extract_xlsx_lines(file_path: str) -> list[str]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    lines = [f"Sheet: {worksheet.title}"]

    for row in worksheet.iter_rows(min_row=1, max_row=8, values_only=True):
        values = [normalize_spaces(str(value)) for value in row if value not in (None, "")]
        if not values:
            continue
        lines.append(" | ".join(values[:4]))
        if len(lines) >= 8:
            break

    return lines or ["Spreadsheet template", "Preview values were not available."]


def get_file_preview_lines(file_path: str) -> list[str]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_lines(file_path)
    if suffix == ".xlsx":
        return extract_xlsx_lines(file_path)
    return ["Digital download", Path(file_path).name]


def draw_gradient(image: Image.Image, top_color, bottom_color) -> None:
    width, height = image.size
    draw = ImageDraw.Draw(image)
    for y in range(height):
        ratio = y / max(height - 1, 1)
        color = tuple(int(top_color[i] * (1 - ratio) + bottom_color[i] * ratio) for i in range(3))
        draw.line((0, y, width, y), fill=color)


def draw_fitted_title(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, max_width: int, max_lines: int = 3) -> int:
    for font in (TITLE_FONT, TITLE_FONT_MEDIUM, TITLE_FONT_SMALL):
        lines = wrap_text(draw, text, font, max_width)
        if len(lines) <= max_lines:
            current_y = y
            line_height = 82 if font == TITLE_FONT else 74 if font == TITLE_FONT_MEDIUM else 66
            for line in lines:
                draw.text((x, current_y), line, font=font, fill=(30, 23, 18))
                current_y += line_height
            return current_y

    lines = wrap_text(draw, text, TITLE_FONT_SMALL, max_width)[:max_lines]
    current_y = y
    for line in lines:
        draw.text((x, current_y), line, font=TITLE_FONT_SMALL, fill=(30, 23, 18))
        current_y += 66
    return current_y


def draw_single_preview(entry: dict, output_path: Path) -> None:
    style = CATEGORY_STYLES.get(entry["category"], CATEGORY_STYLES["Business"])
    image = Image.new("RGB", CANVAS_SIZE, style["bg_top"])
    draw_gradient(image, style["bg_top"], style["bg_bottom"])
    draw = ImageDraw.Draw(image)

    draw.rounded_rectangle((70, 70, 1330, 1680), radius=56, fill=style["panel"], outline=(255, 255, 255), width=3)
    draw.rounded_rectangle((110, 110, 360, 160), radius=24, fill=style["accent"], outline=None)
    draw.text((138, 121), entry["category"].upper(), font=LABEL_FONT, fill=(255, 247, 240))

    title_bottom = draw_fitted_title(draw, entry["name"], 110, 215, 1120)
    draw.text((114, max(title_bottom + 22, 410)), entry["subcategory"], font=SUBTITLE_FONT, fill=style["accent"])

    draw.rounded_rectangle((110, 470, 1290, 1225), radius=42, fill=(255, 255, 255), outline=style["accent_soft"], width=3)

    lines = get_file_preview_lines(entry["files"][0]["localPath"])
    y = 545
    max_width = 1100
    for index, line in enumerate(lines[:7]):
        font = SUBTITLE_FONT if index == 0 else BODY_FONT
        wrapped = wrap_text(draw, line, font, max_width)
        for chunk in wrapped[:2]:
            draw.text((150, y), chunk, font=font, fill=(53, 43, 36))
            y += 58 if font == SUBTITLE_FONT else 48
        y += 10
        if y > 1110:
            break

    draw.rounded_rectangle((110, 1275, 1290, 1505), radius=38, fill=(248, 243, 237), outline=style["accent_soft"], width=2)
    draw.text((150, 1325), "WHAT BUYERS GET", font=LABEL_FONT, fill=style["accent"])
    summary_lines = wrap_text(draw, entry["summary"], SMALL_FONT, 1020)
    summary_y = 1380
    for line in summary_lines[:4]:
        draw.text((150, summary_y), line, font=SMALL_FONT, fill=(92, 75, 64))
        summary_y += 36

    footer_text = f"{entry['productType']}  •  {entry['status']}"
    draw.text((110, 1575), footer_text, font=SMALL_FONT, fill=style["accent"])
    draw.text((110, 1625), Path(entry["files"][0]["storagePath"]).name, font=MONO_FONT, fill=(108, 91, 78))

    image.save(output_path)


def draw_bundle_preview(entry: dict, output_path: Path) -> None:
    style = CATEGORY_STYLES.get(entry["category"], CATEGORY_STYLES["Business"])
    image = Image.new("RGB", CANVAS_SIZE, style["bg_top"])
    draw_gradient(image, style["bg_top"], style["bg_bottom"])
    draw = ImageDraw.Draw(image)

    draw.rounded_rectangle((70, 70, 1330, 1680), radius=56, fill=style["panel"], outline=(255, 255, 255), width=3)
    draw.rounded_rectangle((110, 110, 460, 160), radius=24, fill=style["accent"], outline=None)
    draw.text((142, 121), f"{len(entry['files'])} FILE BUNDLE", font=LABEL_FONT, fill=(255, 247, 240))

    title_bottom = draw_fitted_title(draw, entry["name"], 110, 215, 1120)
    draw.text((114, max(title_bottom + 22, 410)), entry["subcategory"], font=SUBTITLE_FONT, fill=style["accent"])

    card_specs = [
        ((170, 560, 760, 1160), -6),
        ((420, 610, 1010, 1210), 5),
        ((670, 660, 1260, 1260), -4),
    ]

    file_groups = [entry["files"][0:1], entry["files"][1:2], entry["files"][2:3]]
    for (box, _rotation), file_group in zip(card_specs, file_groups):
        x1, y1, x2, y2 = box
        draw.rounded_rectangle(box, radius=34, fill=(255, 255, 255), outline=style["accent_soft"], width=3)
        preview_lines = get_file_preview_lines(file_group[0]["localPath"]) if file_group else ["Bundle file", ""]
        draw.text((x1 + 34, y1 + 34), Path(file_group[0]["storagePath"]).stem.replace("_", " ")[:32], font=LABEL_FONT, fill=style["accent"])
        preview_y = y1 + 92
        for line in preview_lines[:4]:
            for chunk in wrap_text(draw, line, SMALL_FONT, (x2 - x1) - 68)[:2]:
                draw.text((x1 + 34, preview_y), chunk, font=SMALL_FONT, fill=(70, 58, 49))
                preview_y += 34
            preview_y += 10

    draw.rounded_rectangle((110, 1320, 1290, 1530), radius=38, fill=(248, 243, 237), outline=style["accent_soft"], width=2)
    draw.text((150, 1365), "INSIDE THIS BUNDLE", font=LABEL_FONT, fill=style["accent"])
    item_lines = [Path(file["storagePath"]).stem.replace("_", " ") for file in entry["files"][:6]]
    listing = " • ".join(item_lines)
    wrapped_listing = wrap_text(draw, listing, SMALL_FONT, 1000)
    listing_y = 1418
    for line in wrapped_listing[:3]:
        draw.text((150, listing_y), line, font=SMALL_FONT, fill=(92, 75, 64))
        listing_y += 36

    draw.text((110, 1590), entry["summary"], font=SMALL_FONT, fill=(92, 75, 64))
    draw.text((110, 1640), f"{entry['productType']}  •  {entry['status']}", font=SMALL_FONT, fill=style["accent"])

    image.save(output_path)


def main() -> int:
    if len(sys.argv) < 3:
      print("Usage: generate-product-previews.py <manifest.json> <output-dir>")
      return 1

    manifest_path = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    output_dir.mkdir(parents=True, exist_ok=True)

    entries = json.loads(manifest_path.read_text(encoding="utf-8"))

    for entry in entries:
        output_path = output_dir / f"{entry['slug']}.png"
        if len(entry["files"]) > 1:
            draw_bundle_preview(entry, output_path)
        else:
            draw_single_preview(entry, output_path)

    print(f"Generated {len(entries)} preview images in {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
